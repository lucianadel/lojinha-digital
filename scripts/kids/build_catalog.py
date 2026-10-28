"""Prepare local draft metadata, cover artwork and an XLSX review. No network writes.
Only the first (cover) page is rendered. Complete paid PDFs remain in their source.
"""
import argparse
import base64
import csv
import hashlib
import html
import json
import re
import textwrap
from collections import Counter
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

from inspect_acervo import canonical

CATEGORIES = ["Alfabetização", "Matemática", "Raciocínio e Coordenação", "Datas Comemorativas", "Desenhos para Colorir", "Jogos e Recortes"]
COLORS = ["#7bb5ff", "#ffd06d", "#89dfc9", "#ffabc9", "#b9a4ff", "#ffb47d"]
SUMMARIES = {
    "Alfabeto e Letras": "Traçados de letras e propostas de registro com pontilhados para explorar o alfabeto no papel.",
    "Alfabeto Ilustrado e Fichas": "Fichas ilustradas para pintar, recortar e utilizar nas propostas de jogo e estudo do material.",
    "Caça-Letras": "Propostas para localizar letras em conjuntos visuais, pintar as ocorrências e registrar a contagem solicitada.",
    "Caça-Palavras Ilustrados": "Caça-palavras com referências ilustradas e orientação para encontrar palavras na horizontal e na vertical.",
    "Cruzadinhas Ilustradas": "Cruzadinhas em que as figuras ajudam a identificar as palavras e preencher os espaços correspondentes.",
    "Escreva o Nome e Ordem Alfabética": "Atividades de escrita de nomes de figuras e organização de palavras em ordem alfabética.",
    "Famílias Silábicas": "Propostas de leitura de sílabas e traçado sobre letras cinzas, organizadas em famílias silábicas.",
    "Letra Cursiva": "Folhas de treino de escrita cursiva, com modelos em cinza e espaço para continuar o registro.",
    "Letras e Vogais": "Atividades para completar palavras com letras e vogais nas propostas apresentadas no PDF.",
    "Ligue a Figura ao Nome": "Propostas de associação entre figuras e palavras, ligando cada imagem ao nome correspondente.",
    "Sílabas 1: Complete e Forme Palavras": "Atividades para completar palavras com apoio de um banco de sílabas.",
    "Sílabas 2: Separe e Conte": "Propostas para separar palavras em sílabas, escrever nos espaços e contar as partes.",
    "Treino de Escrita: 258 Palavras": "Treinos que combinam figura, leitura de palavra, pontilhado e espaço para escrita independente.",
    "Adição Armada": "Contas de adição organizadas na forma armada, com espaço para registrar os resultados.",
    "Adição com Figuras": "Propostas de adição com apoio de figuras e registro do resultado nos espaços indicados.",
    "Contas com Reserva e 3 Algarismos": "Atividades para armar e resolver contas com reserva e números de três algarismos.",
    "Conte e Escreva": "Atividades de contagem de conjuntos de figuras e escrita da quantidade correspondente.",
    "Divisão": "Propostas de divisão com espaço para resolução e registro dos resultados.",
    "Frações": "Atividades de observação de partes pintadas e registro da fração correspondente.",
    "Hora Certa: Relógio": "Propostas para desenhar ponteiros e representar nos relógios as horas solicitadas.",
    "Números de 0 a 20": "Atividades com numerais de 0 a 20, traçados, contagem de figuras e pintura de quantidades.",
    "Números por Extenso": "Exercícios para observar os números e escrever suas representações por extenso.",
    "Probleminhas de Matemática": "Problemas de matemática com enunciados, espaço para fazer a conta e registrar a resposta.",
    "Sequências e Quadro até 100": "Propostas de completar sequências numéricas e trabalhar com o quadro de números até 100.",
    "Sistema de Numeração": "Atividades de contagem de dezenas e unidades com representação por barrinhas e cubinhos.",
    "Subtração Armada": "Contas de subtração organizadas na forma armada, com espaço para registrar os resultados.",
    "Subtração com Figuras": "Propostas de subtração apoiadas em figuras e registro dos resultados nos espaços indicados.",
    "Tabuada e Multiplicação": "Atividades para completar tabuadas e trabalhar com multiplicação nas propostas do PDF.",
    "Vizinhos, Maior e Menor": "Atividades com números vizinhos e relações de maior e menor, incluindo antecessor e sucessor.",
    "Atenção e Percepção": "Propostas de observar fileiras e identificar a figura diferente em cada conjunto.",
    "Coordenação Motora": "Percursos pontilhados para traçar com lápis e levar os personagens até o destino indicado.",
    "Copie o Desenho (Quadriculado)": "Propostas de observar um modelo e reproduzir seu desenho pintando uma grade quadriculada.",
    "Ligue os Pontos": "Atividades de ligar pontos numerados para revelar desenhos e depois colori-los.",
    "Pinte pelo Código": "Propostas de pintura de quadrinhos seguindo uma correspondência entre números e cores.",
    "Sequência Lógica": "Atividades de observar padrões e completar sequências com a figura correspondente.",
    "Sudoku Infantil": "Grades de sudoku com figuras ou números, seguindo as orientações de linhas, colunas e quadrados do PDF.",
    "Bingo de Letras": "Cartelas de bingo de letras para marcar os elementos sorteados nas dinâmicas propostas.",
    "Bingo de Números": "Cartelas de bingo de números para acompanhar o sorteio e marcar os elementos correspondentes.",
    "Bingo de Sílabas": "Cartelas de bingo de sílabas para marcar os elementos sorteados nas dinâmicas do material.",
    "Bingo de Figuras": "Cartelas ilustradas de bingo para reconhecer e marcar as figuras sorteadas.",
    "Jogo da Memória": "Cartas para pintar, recortar e utilizar em jogos de encontrar pares.",
}


def key(value):
    return canonical(value).replace(" ", "")


def theme_from_source(title):
    return title.removeprefix("Kit ")


def summary_for(kit):
    theme = theme_from_source(kit["source_title"])
    if theme.startswith("Colorir:"):
        return f"Desenhos sobre {theme.split(':', 1)[1].strip().lower()} para colorir e completar a proposta de escrita do nome de cada figura."
    if kit["category"] == "Datas Comemorativas":
        return f"Atividades com a temática {theme}, incluindo propostas de contagem de figuras e registro de quantidades presentes no PDF."
    if theme.startswith("Labirintos:"):
        return f"Percursos de labirintos identificados no arquivo como {theme.split(':', 1)[1].strip().lower()}, para traçar o caminho até o destino."
    return SUMMARIES[theme]


def svg_cover(kit, title, color, image_data):
    theme = theme_from_source(kit["source_title"]).replace("Colorir: ", "")
    lines = textwrap.wrap(theme, width=26, break_long_words=False)
    assert len(lines) <= 3
    text = "".join(f'<text x="55" y="{267 + n * 66}" font-size="55" font-weight="700" fill="#fff">{html.escape(line)}</text>' for n, line in enumerate(lines))
    count = f'{kit["index_activities"]:,}'.replace(",", ".")
    return f'''<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="900" height="1200" viewBox="0 0 900 1200">
<title>{html.escape(title)}</title><desc>Capa da coleção Lojinha Digital Kids, com ilustrações extraídas somente da capa original licenciada.</desc>
<rect width="900" height="1200" rx="30" fill="#10233f"/>
<g font-family="Arial, sans-serif"><rect x="55" y="46" width="62" height="62" rx="17" fill="#315de4"/>
<path d="M86 62c-9-6-16-6-23-4v31c9-2 16-1 23 4m0-31c9-6 16-6 23-4v31c-9-2-16-1-23 4V62" fill="none" stroke="#fff" stroke-width="3"/>
<text x="135" y="78" font-size="30" font-weight="700" fill="#fff">LOJINHA DIGITAL</text>
<text x="630" y="86" font-size="47" font-weight="700" fill="{color}">KIDS</text>
<text x="57" y="135" font-size="18" letter-spacing="2" fill="#bdcee9">APRENDER, COLORIR E DESCOBRIR</text>
<rect x="55" y="168" width="790" height="42" rx="20" fill="#24416a"/><text x="76" y="196" font-size="18" font-weight="700" fill="{color}">{html.escape(kit["category"].upper())}</text>
{text}
<rect x="55" y="443" width="790" height="90" rx="22" fill="{color}"/>
<text x="81" y="501" font-size="44" font-weight="700" fill="#10233f">{count} ATIVIDADES</text>
<circle cx="796" cy="487" r="18" fill="#fff" fill-opacity="0.7"/><path d="m786 487 7 7 14-15" fill="none" stroke="#10233f" stroke-width="4"/>
<rect x="45" y="560" width="810" height="395" rx="28" fill="#fff"/>
<image x="65" y="577" width="770" height="360" preserveAspectRatio="xMidYMid meet" xlink:href="data:image/png;base64,{image_data}"/>
<rect x="55" y="985" width="790" height="57" rx="16" fill="#1c3659"/>
<text x="78" y="1023" font-size="25" fill="#e6efff">Faixa indicada no arquivo: {html.escape(kit["age"])}</text>
<text x="55" y="1094" font-size="25" font-weight="700" fill="#fff">PDF • A4 • PARA IMPRIMIR</text>
<text x="55" y="1134" font-size="18" fill="#bdcee9">Atividades em preto e branco • Conteúdo licenciado, não exclusivo</text>
<circle cx="810" cy="1092" r="12" fill="#ed80b5"/><circle cx="778" cy="1126" r="8" fill="#7bb5ff"/>
</g></svg>'''


def build(audit_path, project_root):
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["summary"]["all_counts_match"] and audit["summary"]["activities_total"] == 31693
    output = project_root / "kids-review"
    (output / "covers").mkdir(parents=True, exist_ok=True)
    snapshot = json.loads((output / "supabase-public-snapshot.json").read_text(encoding="utf-8"))
    categories = [{"name": name, "slug": canonical(name).replace(" ", "-"), "id": canonical(name).replace(" ", "-")} for name in CATEGORIES]
    products = []
    for kit in sorted(audit["kits"], key=lambda k: (CATEGORIES.index(k["category"]), k["source_title"])):
        theme = theme_from_source(kit["source_title"])
        professional_theme = theme.split(": ", 1)[1] + " para Colorir" if theme.startswith("Colorir: ") else theme
        n = kit["index_activities"]
        title = f"Kit {professional_theme} — {n} atividades"
        slug = "kids-" + canonical(professional_theme).replace(" ", "-")
        category = next(c for c in categories if c["name"] == kit["category"])
        match_keys = sorted({key(kit["source_title"]), key(title), key(slug), key("Kit " + professional_theme)})
        matches = [p for p in snapshot["products"] if key(p["title"]) in match_keys or key(p["slug"]) in match_keys]
        assert len(matches) <= 1, f"Ambiguous public match for {title}"
        cents = 990 if n <= 80 else 1490 if n <= 200 else 1990
        summary = summary_for(kit)
        short = f"{n} atividades para imprimir em A4. {professional_theme}. Indicação do arquivo: {kit['age']}."
        description = f"""## {professional_theme}
{summary}

### O que está incluído
- {n} páginas de atividades e 1 capa ({kit['pdf_pages']} páginas no arquivo).
- Faixa etária indicada no PDF original: {kit['age']}.
- Arquivo PDF, folhas em preto e branco e tamanho A4, prontas para imprimir.

### Como utilizar
Abra o PDF em um leitor compatível e imprima as páginas que desejar utilizar. Crianças podem contar com a orientação de pais, professores ou responsáveis, especialmente nas propostas de recorte. O produto é digital; a compra não inclui impressão física.

### Compra e entrega
Solicite a compra pelo WhatsApp da loja. O pagamento é combinado no atendimento e o PDF é enviado manualmente após a confirmação do pagamento.

Material de acervo licenciado e não exclusivo. A apresentação desta coleção utiliza a identidade Lojinha Digital Kids."""
        assert len(title) <= 120 and len(short) <= 220 and len(description) <= 10000
        with pymupdf.open(Path(audit["source_root"]) / kit["source_relative_path"]) as doc:
            page = doc[0]
            clip = pymupdf.Rect(25, page.rect.height * .44, page.rect.width - 25, page.rect.height * .825)
            crop = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=clip, alpha=False).tobytes("png")
        svg = svg_cover(kit, title, COLORS[CATEGORIES.index(kit["category"])], base64.b64encode(crop).decode())
        svg_path = output / "covers" / (slug + ".svg")
        svg_path.write_text(svg, encoding="utf-8")
        with pymupdf.open(stream=svg.encode(), filetype="svg") as cover:
            cover[0].get_pixmap(matrix=pymupdf.Matrix(4/3, 4/3), alpha=False).save(output / "covers" / (slug + ".png"))
        notes = []
        if kit["alternate_files"]:
            notes.append("Há versão avulsa diferente: esta prévia usa o PDF da pasta extraída do ZIP; confirmar a escolha.")
        if "Corpo Humano" in title:
            notes.append("Conferir cadastro também entre rascunhos. Se já estiver a R$ 5,90, manter o registro e o preço; não sobrescrever.")
        product = {"title": title, "slug": slug, "category_name": category["name"], "category_slug": category["slug"],
                   "activities": n, "age": kit["age"], "pdf_pages": kit["pdf_pages"], "short_description": short,
                   "description": description, "price": cents / 100, "price_cents": cents,
                   "cover_file": "covers/" + slug + ".png", "cover_svg": "covers/" + slug + ".svg", "cover_url": None,
                   "cover_alt": f"Capa Lojinha Digital Kids: {professional_theme}, {n} atividades, indicação de {kit['age']}.",
                   "purchase_method": "whatsapp", "whatsapp_number": "5521992530081", "hotmart_url": None, "status": "draft", "featured": False,
                   "source_title": kit["source_title"], "source_relative_path": kit["source_relative_path"], "source_sha256": kit["sha256"],
                   "match_keys": match_keys, "known_existing_product": matches[0] if matches else None,
                   "review_keys": [key(theme.split(": ", 1)[1] if theme.startswith("Colorir: ") else theme)] if len(key(theme.split(": ", 1)[1] if theme.startswith("Colorir: ") else theme)) >= 6 else [],
                   "review_notes": notes, "content_summary": summary, "activity_count_verified": all(kit["checks"].values())}
        products.append(product)
    assert len(products) == 63 and len({p["slug"] for p in products}) == 63
    def select(source_title):
        return next(p for p in products if p["source_title"] == source_title)
    combos = []
    combo_specs = [
        ("Combo Colorir e Aprender", ["Kit Colorir: Corpo Humano", "Kit Colorir: Natureza", "Kit Colorir: Animais"], 2490),
        ("Combo Primeiros Passos", ["Kit Alfabeto e Letras", "Kit Números de 0 a 20", "Kit Coordenação Motora"], 2990),
        ("Combo Professor Criativo", ["Kit Famílias Silábicas", "Kit Adição com Figuras", "Kit Jogo da Memória"], 2990),
    ]
    for title, names, cents in combo_specs:
        members = [select(name) for name in names]
        total = sum(p["price_cents"] for p in members)
        combos.append({"title": title, "source_titles": names, "slugs": [p["slug"] for p in members], "activities": sum(p["activities"] for p in members),
                       "suggested_price": cents / 100, "individual_sum": total / 100, "guide_reference": 29.9,
                       "notes": "A referência de R$ 29,90 supera a soma de R$ 29,70 dos avulsos; R$ 24,90 é uma alternativa promocional, sujeita à aprovação. Se Corpo Humano já custar R$ 5,90, recalcular a economia." if cents == 2490 else "Proposta de três kits; valores, composição e versões dependem de aprovação.",
                       "status": "proposal_only"})
    for category in categories:
        members = [p for p in products if p["category_slug"] == category["slug"]]
        combos.append({"title": "Bloco completo — " + category["name"], "source_titles": [p["source_title"] for p in members], "slugs": [p["slug"] for p in members],
                       "activities": sum(p["activities"] for p in members), "suggested_price": 47.0,
                       "individual_sum": sum(p["price_cents"] for p in members) / 100, "guide_reference": 47.0,
                       "notes": "Bloco completo proposto, sem criação de cadastro ou união automática dos PDFs.", "status": "proposal_only"})
    for category in categories:
        aliases = {key(category["name"]), key(category["slug"])}
        category["match_keys"] = sorted(aliases)
        category["known_matches"] = [c for c in snapshot["categories"] if key(c["name"]) in aliases or key(c["slug"]) in aliases]
        assert len(category["known_matches"]) <= 1
        category["kit_count"] = sum(p["category_slug"] == category["slug"] for p in products)
        category["activities"] = sum(p["activities"] for p in products if p["category_slug"] == category["slug"])
    bundle = {"collection": "Lojinha Digital Kids — Aprender, Colorir e Descobrir", "status": "proposal_for_owner_approval", "categories": categories,
              "products": products, "combos": combos, "summary": audit["summary"],
              "pricing_source": "Faixas fornecidas pela proprietária na solicitação. Guia da Primeira Venda separado não localizado.",
              "pricing_boundary": "Até 80 inclui 80; a faixa seguinte começa em 81.",
              "existing_records_visibility": "Somente catálogo público. A conferência definitiva deve incluir rascunhos sob permissão administrativa antes de importar.",
              "source_index_sha256": audit["index_sha256"]}
    (output / "catalogo-proposta.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    template = (project_root / "scripts/kids/review-template.html").read_text(encoding="utf-8")
    visible_fields = {"title", "slug", "category_name", "category_slug", "activities", "age", "pdf_pages", "price", "cover_file", "cover_alt", "content_summary", "description", "review_notes"}
    review_data = {"categories": categories, "products": [{key: value for key, value in p.items() if key in visible_fields} for p in products], "combos": combos}
    embedded = json.dumps(review_data, ensure_ascii=False).replace("<", "\\u003c")
    (output / "index.html").write_text(template.replace("__KIDS_REVIEW_DATA__", embedded), encoding="utf-8")
    with (output / "catalogo-63-kits.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        fields = ["title", "slug", "category_name", "activities", "age", "pdf_pages", "price", "cover_file", "cover_alt", "short_description", "description", "purchase_method", "whatsapp_number", "status", "source_title", "source_sha256", "review_notes"]
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows({**p, "review_notes": " | ".join(p["review_notes"])} for p in products)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Leia primeiro"
    for row in [["LOJINHA DIGITAL KIDS", "REVISÃO — NÃO PUBLICADO"], ["Kits", 63], ["Atividades", 31693], ["PDFs encontrados", 65], ["Versões alternativas", 2],
                ["Preço", "Sugestões sujeitas à aprovação; não são preços finais."], ["Fontes", "Índice/licença e cada PDF original. PDF pago não acompanha esta planilha."],
                ["Rascunhos existentes", "Não visíveis na consulta pública. Conferir antes da importação."], ["Corpo Humano", "8 atividades; manter R$ 5,90 se já cadastrado assim, sem sobrescrever."],
                ["Direitos", "Conteúdo licenciado e não exclusivo. Preservar créditos e restrições da licença."], ["Guia de preços", bundle["pricing_source"]]]:
        summary_sheet.append(row)
    sheet = workbook.create_sheet("63 kits")
    fields = ["title", "slug", "category_name", "activities", "age", "pdf_pages", "price", "status", "purchase_method", "whatsapp_number", "cover_file", "cover_alt", "short_description", "description", "source_title", "source_relative_path", "source_sha256", "review_notes"]
    sheet.append(["Título", "Slug", "Categoria", "Atividades", "Idade indicada", "Páginas PDF", "Preço sugerido", "Status inicial", "Compra", "WhatsApp", "Capa", "Texto alternativo", "Descrição curta", "Descrição completa", "Kit original", "Arquivo original (relativo)", "SHA-256", "Pontos para aprovação"])
    for p in products:
        sheet.append([" | ".join(p[f]) if isinstance(p[f], list) else p[f] for f in fields])
        sheet.cell(sheet.max_row, 7).number_format = '"R$" #,##0.00'
        sheet.cell(sheet.max_row, 11).hyperlink = p["cover_file"]
    categories_sheet = workbook.create_sheet("6 categorias")
    categories_sheet.append(["Categoria", "Slug", "Kits", "Atividades", "Categoria pública equivalente"])
    for c in categories:
        categories_sheet.append([c["name"], c["slug"], c["kit_count"], c["activities"], " | ".join(m["name"] for m in c["known_matches"]) or "Nenhuma identificada; criar somente após aprovação"])
    combo_sheet = workbook.create_sheet("Combos e blocos")
    combo_sheet.append(["Proposta", "Kits incluídos", "Atividades", "Soma avulsos", "Preço sugerido", "Referência guia", "Observações"])
    for c in combos:
        combo_sheet.append([c["title"], "\n".join(c["source_titles"]), c["activities"], c["individual_sum"], c["suggested_price"], c["guide_reference"], c["notes"]])
        for col in [4, 5, 6]: combo_sheet.cell(combo_sheet.max_row, col).number_format = '"R$" #,##0.00'
    cover_sheet = workbook.create_sheet("Capas")
    cover_sheet.append(["Kit", "Prévia", "Atividades", "Idade indicada", "Arquivo da capa"])
    for p in products:
        cover_sheet.append([p["title"], "", p["activities"], p["age"], p["cover_file"]])
        image = Image(output / p["cover_file"])
        image.width, image.height = 110, 147
        cover_sheet.add_image(image, f"B{cover_sheet.max_row}")
        cover_sheet.row_dimensions[cover_sheet.max_row].height = 116
    for ws in workbook:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="10233F")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0: cell.fill = PatternFill("solid", fgColor="EDF3FF")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 30 if col[0].column > 1 else 55
    sheet.column_dimensions["N"].width = 70
    cover_sheet.column_dimensions["B"].width = 19
    workbook.save(output / "Lojinha-Digital-Kids-63-kits.xlsx")
    # Lightweight app lookup: no paid content, proposed prices, local paths or PDF URLs.
    app_data = {"categories": [{"id": c["id"], "name": c["name"], "slug": c["slug"]} for c in categories],
                "kits": [{"slug": p["slug"], "sourceTitle": p["source_title"], "categorySlug": p["category_slug"], "matchKeys": p["match_keys"]} for p in products]}
    data_path = project_root / "src/data/kidsCollection.json"
    data_path.write_text(json.dumps(app_data, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "creditos-e-licenca.md").write_text("# Referência de licença\n\nFonte: 00 - LEIA-ME - Índice e Licença.pdf, páginas 3 e 4.\n\nA licença informa revenda PLR não exclusiva e permite capas, nomes e marca próprios. Proíbe redistribuir o acervo completo com direitos de revenda, liberar o acervo completo gratuitamente e reivindicar exclusividade sobre o conteúdo. Os créditos originais citam Noto Emoji/Google (Apache 2.0) e as fontes Andika, Fredoka, Nunito e Playwrite BR (SIL OFL 1.1). Mantenha os avisos/créditos exigidos pelo documento; esta coleção não atribui direitos de revenda aos clientes.\n\nAs ilustrações das capas de revisão foram extraídas exclusivamente da primeira página dos PDFs originais. Nenhuma imagem externa foi adicionada e nenhuma página de atividade foi publicada como amostra.\n", encoding="utf-8")
    print(json.dumps({"kits": len(products), "activities": sum(p["activities"] for p in products), "price_counts": dict(Counter(str(p["price"]) for p in products)), "covers_png": len(list((output / 'covers').glob('*.png'))), "known_existing_products": sum(bool(p['known_existing_product']) for p in products), "combos_and_blocks": len(combos)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    args = parser.parse_args()
    build(args.audit, args.project_root.resolve())
